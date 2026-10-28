"""Synthetic inputs for the app tests (this file is not itself a test module).

  make_pptx_bytes / make_pdf_bytes        small slide files with known content
  make_profile                            a complete synthetic handwriting profile (solid-block glyphs)
  make_completed_template_bytes           the real blank template with "handwriting" drawn in every box
"""

import io
import json
import os
import string
import sys
import tempfile
from pathlib import Path

import pymupdf
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import config  # noqa: E402
import spacing  # noqa: E402
import template  # noqa: E402

ALL_CHARS = string.ascii_uppercase + string.ascii_lowercase + string.digits
LETTERS_AND_DIGITS = ALL_CHARS
V2_CHARS = ALL_CHARS + config.PUNCTUATION_CHARS


def make_pptx_bytes(with_picture_slide=False, title="Course Overview"):
    """A deck with a title slide, a bullet slide with nesting, and optionally a picture-only slide."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = "Spring 2025"
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Main Topic"
    body = slide.placeholders[1].text_frame
    body.text = "First point"
    for text, level in (("Second point, with punctuation!", 0), ("Supporting detail: 3.5% on 2025-03-14", 1)):
        paragraph = body.add_paragraph()
        paragraph.text = text
        paragraph.level = level
    if with_picture_slide:
        picture = io.BytesIO()
        Image.new("RGB", (40, 40), (200, 100, 50)).save(picture, "PNG")
        picture.seek(0)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(picture, Inches(1), Inches(1))
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def make_pdf_bytes(pages=2):
    """A text-based slide PDF: each page has a large title and three bullets."""
    document = pymupdf.open()
    for number in range(1, pages + 1):
        page = document.new_page(width=960, height=540)
        page.insert_text((60, 90), f"Slide Title {number}", fontsize=36, fontname="hebo")
        for i, text in enumerate((f"First bullet of page {number}", "Second bullet, with punctuation!", "Third bullet 2025")):
            page.insert_text((60, 170 + 40 * i), "•", fontsize=20, fontname="helv")
            page.insert_text((82, 170 + 40 * i), text, fontsize=20, fontname="helv")
    return document.tobytes()


def make_image_only_pdf_bytes():
    document = pymupdf.open()
    page = document.new_page(width=960, height=540)
    page.draw_rect(pymupdf.Rect(0, 0, 960, 540), fill=(0.8, 0.7, 0.5))
    return document.tobytes()


def make_profile(root, name, characters=ALL_CHARS, style=True, letter_gap=2.0, word_gap=12.0, template_version=None,
                 baselines=None):
    """A synthetic handwriting profile folder: solid black blocks, 3 samples per character."""
    folder = Path(root) / name
    folder.mkdir(parents=True, exist_ok=True)
    records = []
    for index, char in enumerate(characters):
        for sample in (1, 2, 3):
            width, height = 24 + 2 * sample + (6 if char.isupper() else 0), 46 if char in "gjpqy" else 30
            file_name = f"glyph_{index:03d}_{sample}.png"          # nothing here can be guessed from the character
            Image.new("RGBA", (width, height), (0, 0, 0, 255)).save(folder / file_name)
            records.append({"char": char, "sample": sample, "file": file_name, "baseline_from_top_px": (baselines or {}).get(char, 30.0),
                            "image_size_px": {"width": width, "height": height}})
    (folder / config.GLYPH_METADATA_NAME).write_text(json.dumps(
        {"meta": {"pixels_per_point": 4.0, "trim_padding_px": 4,
                  **({"template_version": template_version} if template_version else {})}, "glyphs": records}))
    if style:
        (folder / config.STYLE_NAME).write_text(json.dumps(
            {"letter_spacing": {"estimate_pt": letter_gap, "quality": "approximate"},
             "word_spacing": {"mean_pt": word_gap, "quality": "good"}}))
    return folder


def make_completed_template_bytes(skip=lambda box: False, sentence=True, pages=None, version=None):
    """(Newest template version unless `version` is given.) The blank template PDF with a black stem drawn in every box and the sentence 'written' in its area."""
    layout = template.build_layout(version)
    with tempfile.TemporaryDirectory() as folder:
        path = os.path.join(folder, "blank.pdf")
        template.write_pdf(layout, path)
        document = pymupdf.open(path)
        height = layout["pages"][0]["height"]
        for box in layout["boxes"]:
            if skip(box):
                continue
            x0 = box["x0"] + 10
            document[box["page"] - 1].draw_rect(
                pymupdf.Rect(x0, height - (box["baseline_y"] + 14), x0 + 3, height - (box["baseline_y"] - 6)),
                color=None, fill=(0, 0, 0))
        if sentence:
            area = layout["sample_sentence"]
            words = spacing.sentence_words(area["text"])
            for line_index, line_words in enumerate((words[:6], words[6:])):
                baseline, x = area["baseline_ys"][line_index], area["x0"] + 10
                for word in line_words:
                    for letter_index in range(len(word)):
                        document[area["page"] - 1].draw_rect(
                            pymupdf.Rect(x, height - baseline - 12, x + 5, height - baseline), color=None, fill=(0, 0, 0))
                        x += 5 + (2 if letter_index < len(word) - 1 else 12)
        if pages is not None:
            document.select(list(range(pages)))
        return document.tobytes()
