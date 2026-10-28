"""Extract readable text from a PowerPoint (.pptx) or text-based PDF, ready for render.py.

Usage:
    python slides.py slides_in/lecture.pptx --output output/lecture_notes.txt
    python render.py --input output/lecture_notes.txt --user amaya

Writes three things (the input file is only ever read):
    <output>.txt     plain notes text that render.py understands
    <output>.json    the same content as a structured document (schema below)
    --preview FILE   optional slide-by-slide structure listing for checking by eye

How the two formats differ
  .pptx   Structure is READ from the file: title placeholders, paragraphs, bullet
          settings (explicit, or inherited from the layout/master), nesting levels,
          auto-numbering. Reading order is by position (rows, then left to right).
  .pdf    A PDF has no titles or bullets, only positioned text. Titles, headings,
          bullets and nesting are INFERRED from font size, position and bullet
          characters (see slides_pdf.py). Every inferred element says so in the JSON.

JSON schema (schema_version 1)
  source        {file, path, format, size_bytes, sha256, pages}
  extracted_at  UTC time
  method        short description of how the text was read
  structure_note  explicit vs inferred structure, in words
  slides[]      {number, title, title_source, status, hidden?, elements[], warnings[]}
      title_source  "placeholder" | "inferred_position" | "inferred_font_size" | "none"
      status        "ok" | "no_text" | "uncertain"
      elements[]    {type, text, level, structure, marker?}
          type       "title" | "subtitle" | "heading" | "paragraph" | "bullet" | "table_row"
          level      nesting depth, 0 = top level
          structure  "explicit" (read from the file) | "inferred" (guessed from layout)
          marker     only for numbered bullets, e.g. "1."
  warnings[]    document-level notes (replaced characters, dropped footers, ...)
  summary       {slides, slides_with_text, elements}

Text conventions (notes .txt)
  * Text is plain ASCII: typographic quotes/dashes become ' " -, accents are removed,
    other symbols become '?' (and are reported in warnings).
  * Each slide starts with its title; a slide with no title gets the heading
    "Slide N" (config.SLIDES_FALLBACK_HEADING). No title text is ever invented.
  * Bullets are "* text", nested levels are indented 3 spaces per level, numbered
    bullets keep their number ("1. text").
  * A blank line separates elements, except between two bullets and between a
    title and its subtitle. Slides are separated by 2 blank lines. Slides without any
    text are left out (listed in warnings).
"""

import argparse
import datetime
import hashlib
import json
import os
import re
import sys
import unicodedata
from pathlib import Path

import config
from extract import ExtractError

A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
BULLET_TAGS = {A_NS + "buNone": "none", A_NS + "buChar": "char", A_NS + "buBlip": "char", A_NS + "buAutoNum": "auto"}
EMU_PER_INCH = 914400


# ---------------------------------------------------------------------------
# Text cleaning
# ---------------------------------------------------------------------------

def clean_text(text):
    """Make text safe for the renderer. Returns (ascii text, set of characters turned into '?')."""
    text = text.replace("\v", " ").replace("\n", " ").replace("\r", " ")
    for old, new in {**config.TYPOGRAPHIC_REPLACEMENTS, **config.SLIDE_TEXT_REPLACEMENTS}.items():
        text = text.replace(old, new)
    out, unknown = [], set()
    for char in text:
        if 32 <= ord(char) < 127:
            out.append(char)
        elif char.isspace():
            out.append(" ")
        else:
            # Accented letters decompose to a plain letter plus combining marks.
            base = "".join(c for c in unicodedata.normalize("NFKD", char) if not unicodedata.combining(c))
            if base and all(32 <= ord(c) < 127 for c in base):
                out.append(base)
            else:
                out.append("?")
                unknown.add(char)
    return " ".join("".join(out).split()), unknown


# ---------------------------------------------------------------------------
# PowerPoint extraction
# ---------------------------------------------------------------------------

def paragraph_bullet(paragraph, shape, slide):
    """Is this paragraph bulleted? Returns ("none" | "char" | "auto", the bullet XML element or None).

    PowerPoint stores bullets in several places, so look in the order PowerPoint does:
    the paragraph itself, the shape's own list style, the layout placeholder, then the
    slide master's body style. Text that is not a body placeholder (plain text boxes,
    titles) has no bullets unless it says so explicitly.
    """
    level_tag = f"{A_NS}lvl{min(paragraph.level, 8) + 1}pPr"

    def find_bullet(properties):
        if properties is not None:
            for child in properties:
                if child.tag in BULLET_TAGS:
                    return BULLET_TAGS[child.tag], child
        return None

    found = find_bullet(paragraph._p.find(f"{A_NS}pPr"))
    if found:
        return found

    def from_list_style(element):
        list_style = element.find(f".//{A_NS}lstStyle") if element is not None else None
        return find_bullet(list_style.find(level_tag)) if list_style is not None else None

    found = from_list_style(shape._element)
    if found:
        return found
    if not shape.is_placeholder:
        return "none", None

    layout = slide.slide_layout
    for placeholder in layout.placeholders:
        if placeholder.placeholder_format.idx == shape.placeholder_format.idx:
            found = from_list_style(placeholder._element)
            if found:
                return found
    kind = shape.placeholder_format.type
    title_like = {"TITLE", "CENTER_TITLE", "SUBTITLE", "VERTICAL_TITLE", "DATE", "FOOTER", "SLIDE_NUMBER"}
    if kind is None or kind.name in title_like:
        return "none", None
    master_ns = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
    body_style = layout.slide_master._element.find(".//p:txStyles/p:bodyStyle", master_ns)
    found = find_bullet(body_style.find(level_tag)) if body_style is not None else None
    return found or ("none", None)


def number_marker(element, counter):
    """Marker text for an auto-numbered paragraph: arabic "1." / "1)" or letters "a." / "A)"."""
    style = element.get("type", "arabicPeriod") if element is not None else "arabicPeriod"
    closing = ")" if "Paren" in style else "."
    if style.startswith("alphaLc") or style.startswith("alphaUc"):
        letter = chr(ord("a") + (counter - 1) % 26)
        return (letter.upper() if style.startswith("alphaUc") else letter) + closing
    return f"{counter}{closing}"


def shape_paragraphs(shape, slide, text_kind):
    """Elements for every non-empty paragraph of one text shape."""
    elements, counters = [], {}
    for paragraph in shape.text_frame.paragraphs:
        text = paragraph.text
        if not text.strip():
            continue
        level = min(paragraph.level, config.SLIDES_MAX_LEVEL)
        if text_kind != "paragraph":                     # title / subtitle placeholders
            elements.append({"type": text_kind, "text": text, "level": 0, "structure": "explicit"})
            continue
        kind, bullet = paragraph_bullet(paragraph, shape, slide)
        for deeper in [k for k in counters if k > level]:
            del counters[deeper]
        if kind == "auto":
            start = int(bullet.get("startAt", 1)) if bullet is not None else 1
            counters[level] = counters.get(level, start - 1) + 1
            elements.append({"type": "bullet", "text": text, "level": level, "structure": "explicit",
                             "marker": number_marker(bullet, counters[level])})
        elif kind == "char":
            counters.pop(level, None)
            elements.append({"type": "bullet", "text": text, "level": level, "structure": "explicit"})
        else:
            counters.pop(level, None)
            elements.append({"type": "paragraph", "text": text, "level": level, "structure": "explicit"})
    return elements


def flatten_shapes(shapes, anchor=None):
    """Yield (anchor position, inner position, shape), looking inside groups."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for shape in shapes:
        position = (shape.top or 0, shape.left or 0)
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from flatten_shapes(shape.shapes, anchor or position)
        else:
            yield (anchor or position), (position if anchor else (0, 0)), shape


def reading_order(entries):
    """Order shapes top to bottom in rows; shapes in the same row go left to right."""
    tolerance = config.PPTX_ROW_TOLERANCE_IN * EMU_PER_INCH
    ordered, row_top, row = [], None, []
    for entry in sorted(entries, key=lambda e: e["anchor"][0]):
        if row_top is None or entry["anchor"][0] - row_top > tolerance:
            ordered.extend(sorted(row, key=lambda e: (e["anchor"][1], e["inner"])))
            row, row_top = [], entry["anchor"][0]
        row.append(entry)
    ordered.extend(sorted(row, key=lambda e: (e["anchor"][1], e["inner"])))
    return ordered


def boxes_overlap(a, b):
    """Fraction of the smaller box covered by the overlap of two (left, top, width, height) boxes."""
    width = min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0])
    height = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
    smaller = min(a[2] * a[3], b[2] * b[3])
    return 0.0 if width <= 0 or height <= 0 or smaller <= 0 else width * height / smaller


def extract_slide(slide, number, slide_height):
    """Raw slide dict (text not yet cleaned) from one PowerPoint slide."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    warnings, skipped, pictures = [], 0, 0
    title, subtitles, entries = None, [], []

    for anchor, inner, shape in flatten_shapes(slide.shapes):
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            pictures += 1
        kind = shape.placeholder_format.type.name if shape.is_placeholder and shape.placeholder_format.type else None
        if kind in ("DATE", "FOOTER", "SLIDE_NUMBER"):
            skipped += 1
            continue
        box = (shape.left or 0, shape.top or 0, shape.width or 0, shape.height or 0)
        # Charts and SmartArt keep their text in separate parts that are not read; say so.
        for uri in (shape._element.xpath(".//a:graphicData/@uri") if hasattr(shape._element, "xpath") else []):
            for suffix, name in (("/chart", "a chart"), ("/diagram", "a SmartArt diagram")):
                message = f"This slide contains {name}; text inside it is not extracted."
                if uri.endswith(suffix) and message not in warnings:
                    warnings.append(message)

        if getattr(shape, "has_table", False) and shape.has_table:
            rows = [" | ".join(c.text.strip() for c in row.cells if c.text.strip()) for row in shape.table.rows]
            elements = [{"type": "table_row", "text": r, "level": 0, "structure": "explicit"} for r in rows if r]
        elif shape.has_text_frame and shape.text_frame.text.strip():
            if kind in ("TITLE", "CENTER_TITLE"):
                title_text = " ".join(p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip())
                if title is None:
                    title = title_text
                    continue
                elements = [{"type": "paragraph", "text": title_text, "level": 0, "structure": "explicit"}]
            elif kind == "SUBTITLE":
                elements = shape_paragraphs(shape, slide, "subtitle")
            else:
                elements = shape_paragraphs(shape, slide, "paragraph")
        else:
            continue
        if elements:
            entries.append({"anchor": anchor, "inner": inner, "box": box, "elements": elements})

    ordered, seen, kept = reading_order(entries), [], []
    for entry in ordered:
        text = " ".join(" ".join(e["text"].lower().split()) for e in entry["elements"])
        if any(text == old_text and boxes_overlap(entry["box"], old_box) >= config.PPTX_DUPLICATE_OVERLAP
               for old_text, old_box in seen):
            warnings.append("A text shape with identical text overlapping another shape was ignored as a duplicate.")
            continue
        seen.append((text, entry["box"]))
        kept.append(entry)

    title_source = "placeholder" if title else "none"
    if title is None and kept:
        # No title placeholder: use a short single-paragraph text near the top, if there is one.
        first = kept[0]
        top_ok = first["anchor"][0] < config.PPTX_TITLE_ZONE * slide_height
        only = first["elements"]
        if (top_ok and len(only) == 1 and only[0]["type"] == "paragraph" and only[0]["level"] == 0
                and len(only[0]["text"].split()) <= config.PPTX_TITLE_MAX_WORDS):
            title, title_source = only[0]["text"], "inferred_position"
            kept = kept[1:]

    elements = [e for entry in kept for e in entry["elements"]]
    status = "ok" if (title or elements) else "no_text"
    if status == "no_text":
        picture_note = f" It contains {pictures} picture(s); text inside images is not extracted." if pictures else ""
        warnings.append("No text found on this slide." + picture_note)
    hidden = slide._element.get("show") == "0"
    if hidden:
        warnings.append("This slide is hidden in the presentation.")
    return {"number": number, "title": title, "title_source": title_source, "status": status,
            "elements": elements, "warnings": warnings, "hidden": hidden, "skipped_footer_shapes": skipped}


def extract_pptx(path):
    from pptx import Presentation
    try:
        presentation = Presentation(str(path))
    except Exception as error:
        raise ExtractError(f"{path} could not be opened as a PowerPoint file ({type(error).__name__}: {error}). "
                           "Is it a valid .pptx? Old .ppt files are not supported.")
    slides = [extract_slide(slide, number, presentation.slide_height)
              for number, slide in enumerate(presentation.slides, start=1)]
    method = "python-pptx: title placeholders, paragraphs, bullet settings (explicit or inherited), levels, tables"
    note = ("Titles, bullets, numbering and nesting levels are read from the presentation's own settings. "
            "Reading order is by position (rows, then left to right), which can differ from the order the author intended.")
    return slides, method, note, []


# ---------------------------------------------------------------------------
# Building the document
# ---------------------------------------------------------------------------

def file_sha256(path):
    digest = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def build_document(path, fmt, raw_slides, method, note, extra_warnings):
    """Clean the raw slides and assemble the JSON document."""
    replaced, warnings, slides = set(), list(extra_warnings), []
    for raw in raw_slides:
        elements, slide_warnings = [], list(raw["warnings"])
        title, unknown = clean_text(raw["title"]) if raw["title"] else ("", set())
        replaced |= unknown
        if title:
            elements.append({"type": "title", "text": title, "level": 0, "structure":
                             "explicit" if raw["title_source"] == "placeholder" else "inferred"})
        for element in raw["elements"]:
            text, unknown = clean_text(element["text"])
            replaced |= unknown
            if not text:
                continue
            item = {**element, "text": text}
            if item["type"] in ("title", "subtitle", "heading"):
                item["level"] = 0
            elements.append(item)   # repeats are only removed upstream, and only when the shapes overlap
        status = raw["status"] if elements else "no_text"
        if elements and raw["status"] == "no_text":
            status = "ok"
        slide = {"number": raw["number"], "title": title or None, "title_source": raw["title_source"] if title else "none",
                 "status": status, "elements": elements, "warnings": slide_warnings}
        if raw.get("hidden"):
            slide["hidden"] = True
        slides.append(slide)

    if replaced:
        warnings.append("Characters with no plain-ASCII equivalent were replaced with '?': "
                        + " ".join(f"U+{ord(c):04X}" for c in sorted(replaced)))
    with_text = [s for s in slides if s["elements"]]
    unreadable = [s["number"] for s in slides if s["status"] != "ok"]
    if unreadable:
        warnings.append("Slides/pages without reliably extracted text: " + ", ".join(map(str, unreadable)))
    return {
        "schema_version": 1,
        "source": {"file": os.path.basename(path), "path": str(path), "format": fmt,
                   "size_bytes": os.path.getsize(path), "sha256": file_sha256(path), "pages": len(slides)},
        "extracted_at": datetime.datetime.now(datetime.timezone.utc).isoformat(timespec="seconds"),
        "method": method,
        "structure_note": note,
        "slides": slides,
        "warnings": warnings,
        "summary": {"slides": len(slides), "slides_with_text": len(with_text),
                    "elements": sum(len(s["elements"]) for s in slides)},
    }


# ---------------------------------------------------------------------------
# Output files
# ---------------------------------------------------------------------------

def element_line(element):
    """The notes line for one element (no trailing newline)."""
    indent = " " * (config.SLIDES_INDENT_SPACES * element["level"])
    if element["type"] == "bullet":
        return f"{indent}{element.get('marker') or config.SLIDES_BULLET_MARKER} {element['text']}"
    return f"{indent}{element['text']}"


def to_notes_text(document):
    """Notes text for render.py: one block per slide, see the module docstring for the rules."""
    blocks = []
    for slide in document["slides"]:
        if not slide["elements"]:
            continue
        elements = list(slide["elements"])
        if not slide["title"]:
            heading = config.SLIDES_FALLBACK_HEADING.format(number=slide["number"])
            elements.insert(0, {"type": "title", "text": heading, "level": 0})
        lines, previous = [], None
        for element in elements:
            if previous is not None:
                both_bullets = previous["type"] == "bullet" and element["type"] == "bullet"
                title_subtitle = previous["type"] == "title" and element["type"] == "subtitle"
                if not (both_bullets or title_subtitle):
                    lines.append("")
            lines.append(element_line(element))
            previous = element
        blocks.append("\n".join(lines))
    return ("\n" * (config.SLIDES_SLIDE_GAP_LINES + 1)).join(blocks) + "\n"


def to_preview(document):
    """Slide-by-slide structure listing for checking the extraction by eye."""
    names = {"title": "TITLE", "subtitle": "SUBTITLE", "heading": "HEADING", "paragraph": "PARA",
             "bullet": "BULLET", "table_row": "TABLE"}
    lines = [f"Source: {document['source']['file']} ({document['source']['format']}, {document['source']['pages']} pages)",
             f"Method: {document['method']}", f"Note:   {document['structure_note']}", ""]
    for slide in document["slides"]:
        lines.append(f"=== {slide['number']}  status={slide['status']}  title={slide['title']!r} ({slide['title_source']})")
        for element in slide["elements"]:
            marks = f"L{element['level']} {element['structure']}"
            lines.append(f"    {names[element['type']]:8s} [{marks}] {element['text'][:110]}")
        for warning in slide["warnings"]:
            lines.append(f"    ! {warning}")
    for warning in document["warnings"]:
        lines.append(f"WARNING: {warning}")
    return "\n".join(lines) + "\n"


def write_text(path, content):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        f.write(content)


# ---------------------------------------------------------------------------
# Command line
# ---------------------------------------------------------------------------

def extract_document(path):
    """Read a .pptx or .pdf and return the structured document."""
    if not os.path.isfile(path):
        raise ExtractError(f"Input file not found: {path}")
    suffix = os.path.splitext(path)[1].lower()
    if suffix == ".pptx":
        raw, method, note, extra = extract_pptx(path)
        fmt = "pptx"
    elif suffix == ".pdf":
        import slides_pdf  # PyMuPDF is only needed for PDFs
        raw, method, note, extra = slides_pdf.extract_pdf(path)
        fmt = "pdf"
    else:
        raise ExtractError(f"Unsupported file type '{suffix or '(none)'}': {path}. Use a .pptx or .pdf file.")
    if not raw:
        raise ExtractError(f"{path} contains no slides or pages.")
    document = build_document(path, fmt, raw, method, note, extra)
    if document["summary"]["elements"] == 0:
        hint = " It may be a scanned or image-only document; OCR is not supported." if fmt == "pdf" else ""
        raise ExtractError(f"No extractable text found in {path}.{hint}")
    return document


def parse_args(argv=None):
    parser = argparse.ArgumentParser(description="Extract slide text for the handwriting renderer.")
    parser.add_argument("input", help="a .pptx or text-based .pdf file")
    parser.add_argument("--output", help="notes text file to write (default: output/<name>_notes.txt)")
    parser.add_argument("--json", help="structured JSON to write (default: same name as the notes file, .json)")
    parser.add_argument("--preview", help="also write a slide-by-slide structure listing to this file")
    return parser.parse_args(argv)


def run(args):
    document = extract_document(args.input)
    stem = os.path.splitext(os.path.basename(args.input))[0]
    output = Path(args.output) if args.output else config.OUTPUT_DIR / f"{stem}_notes.txt"
    json_path = Path(args.json) if args.json else output.with_suffix(".json")
    outputs = [output, json_path] + ([Path(args.preview)] if args.preview else [])
    for path in outputs:
        if path.suffix.lower() in (".pptx", ".pdf", ".ppt") or (path.exists() and os.path.samefile(path, args.input)):
            raise ExtractError(f"Refusing to write to {path}: it would overwrite or look like the input document.")

    write_text(output, to_notes_text(document))
    write_text(json_path, json.dumps(document, indent=2) + "\n")
    if args.preview:
        write_text(args.preview, to_preview(document))

    summary = document["summary"]
    print(f"Extracted {summary['elements']} elements from {summary['slides_with_text']} of {summary['slides']} "
          f"{'slides' if document['source']['format'] == 'pptx' else 'pages'} ({document['source']['format']})")
    for warning in document["warnings"]:
        print(f"WARNING: {warning}")
    print(f"Notes:   {output}")
    print(f"JSON:    {json_path}")
    if args.preview:
        print(f"Preview: {args.preview}")
    return 0


def main(argv=None):
    try:
        return run(parse_args(argv))
    except ExtractError as error:
        print(f"Error: {error}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
